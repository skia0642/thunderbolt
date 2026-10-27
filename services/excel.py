"""회원 엑셀 업로드·다운로드. 다운로드한 파일을 그대로 다시 올려도 데이터가 손실되지 않는다."""
from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from services.attendance import member_attendance_rows
from services.dues import member_dues_rows
from services.errors import ValidationError
from services.members import clean_member, create_member, normalize_phone, update_member
from services.settings import get_settings

COLUMNS = ["이름", "연락처", "포지션", "등번호", "가입일", "상태", "납부방식"]
FIELD = {"이름": "name", "연락처": "phone", "포지션": "position", "등번호": "back_number",
         "가입일": "joined_at", "상태": "status", "납부방식": "pay_type"}
REQUIRED = ("이름", "연락처")
MAX_ROWS = 500

GUIDE = [
    ("열", "필수", "규칙"),
    ("이름", "O", "20자 이하"),
    ("연락처", "O", "010-0000-0000 형식"),
    ("포지션", "", "투수/포수/내야수/외야수 등"),
    ("등번호", "", "0~99 정수, 중복 불가"),
    ("가입일", "", "YYYY-MM-DD"),
    ("상태", "", "활동/휴면 (비우면 신규는 활동, 기존 회원은 현재 값 유지). 휴면은 회비가 부과되지 않습니다."),
    ("납부방식", "", "일괄/분납 (비우면 신규는 일괄, 기존 회원은 현재 값 유지). 분납은 1·4·7·11월 4회."),
    ("", "", "이름+연락처가 같은 기존 회원은 업로드한 내용으로 갱신됩니다."),
]


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "회원"
    ws.append(COLUMNS)
    guide = wb.create_sheet("안내")
    for row in GUIDE:
        guide.append(row)
    for col, w in zip("ABCDEFG", (12, 16, 12, 8, 12, 8, 10)):
        ws.column_dimensions[col].width = w
    guide.column_dimensions["C"].width = 60
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _frame_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
    return buf.getvalue()


def export_members_bytes(conn, year: int | None = None) -> bytes:
    dues = {r["member_id"]: r for r in member_dues_rows(conn, year)}
    att = {r["member_id"]: r for r in member_attendance_rows(conn, year)}
    rows = []
    for m in conn.execute("SELECT * FROM member ORDER BY id"):
        rows.append({
            "이름": m["name"], "연락처": m["phone"], "포지션": m["position"], "등번호": m["back_number"],
            "가입일": m["joined_at"], "상태": m["status"], "납부방식": m["pay_type"],
            "납부액": dues[m["id"]]["paid"], "미납액": dues[m["id"]]["unpaid"],
            "납부율(%)": dues[m["id"]]["rate"], "참석율(%)": att[m["id"]]["rate"],
        })
    return _frame_bytes({"회원": pd.DataFrame(rows, columns=COLUMNS + ["납부액", "미납액", "납부율(%)", "참석율(%)"])})


def export_ledger_bytes(conn, year: int | None = None) -> bytes:
    year = get_settings(conn, year)["year"]
    tx = pd.read_sql_query(
        'SELECT date AS 일자, type AS 구분, category AS 분류, amount AS 금액, memo AS 메모 FROM "transaction" '
        "WHERE strftime('%Y', date) = ? ORDER BY date, id", conn, params=(str(year),))
    return _frame_bytes({"수입지출": tx})


def export_schedule_bytes(conn, year: int | None = None) -> bytes:
    year = str(get_settings(conn, year)["year"])
    sched = pd.read_sql_query(
        "SELECT date AS 날짜, time AS 시간, kind AS 구분, place AS 장소, opponent AS 상대, "
        "venue_fee AS 대관료, status AS 상태, result AS 결과 FROM schedule "
        "WHERE strftime('%Y', date) = ? ORDER BY date", conn, params=(year,))
    att = pd.read_sql_query(
        "SELECT s.date AS 날짜, s.kind AS 구분, m.name AS 이름, "
        "CASE a.attended WHEN 1 THEN '참석' ELSE '결석' END AS 출석 "
        "FROM attendance a JOIN schedule s ON s.id = a.schedule_id JOIN member m ON m.id = a.member_id "
        "WHERE strftime('%Y', s.date) = ? ORDER BY s.date, m.id", conn, params=(year,))
    return _frame_bytes({"일정": sched, "출석": att})


def _cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and pd.isna(v):
        return None
    return v


def parse_members(conn, file) -> list[dict]:
    """엑셀을 읽어 행별 결과를 돌려준다. 반영은 하지 않는다.

    각 항목: {row, action('신규'/'갱신'/'오류'), reason, name, data, member_id}
    """
    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception as e:  # 손상됐거나 xlsx가 아닌 파일
        raise ValidationError("엑셀(.xlsx) 파일을 읽을 수 없습니다. 템플릿 형식을 확인하세요.") from e
    ws = wb["회원"] if "회원" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValidationError("빈 파일입니다.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ValidationError(f"필수 열이 없습니다: {', '.join(missing)}. 템플릿을 사용하세요.")
    if len(rows) - 1 > MAX_ROWS:
        raise ValidationError(f"한 번에 {MAX_ROWS}행까지만 올릴 수 있습니다.")
    idx = {c: header.index(c) for c in COLUMNS if c in header}

    results, seen_keys, seen_backs = [], {}, {}
    for n, row in enumerate(rows[1:], start=2):
        raw = {FIELD[c]: _cell(row[i]) if i < len(row) else None for c, i in idx.items()}
        if all(v is None for v in raw.values()):
            continue
        item = {"row": n, "name": raw.get("name"), "data": None, "member_id": None, "reason": ""}
        try:
            phone = normalize_phone(raw.get("phone"))
            existing = None
            if raw.get("name") and phone:
                existing = conn.execute(
                    "SELECT id, status, pay_type FROM member WHERE name = ? AND phone = ?",
                    (str(raw["name"]).strip(), phone),
                ).fetchone()
            if existing:  # 기존 회원은 상태·납부방식 칸이 비어 있으면 현재 값을 그대로 둔다(기본값으로 덮어쓰지 않는다)
                raw["status"] = raw.get("status") or existing["status"]
                raw["pay_type"] = raw.get("pay_type") or existing["pay_type"]
            # 등번호의 DB 중복은 아래 2차 검증에서 본다(파일 안에서 번호를 서로 맞바꾸는 경우를 허용하기 위해).
            data = clean_member(conn, raw, exclude_id=existing["id"] if existing else None, check_back=False)
            key = (data["name"], data["phone"])
            if key in seen_keys:
                raise ValidationError(f"{seen_keys[key]}행과 이름·연락처가 중복됩니다.")
            back = data["back_number"]
            if back is not None and back in seen_backs:
                raise ValidationError(f"등번호 {back}번이 {seen_backs[back]}행과 중복됩니다.")
            seen_keys[key] = n
            if back is not None:
                seen_backs[back] = n
            item.update(action="갱신" if existing else "신규", data=data, member_id=existing["id"] if existing else None)
        except ValidationError as e:
            item.update(action="오류", reason=" / ".join(e.messages))
        results.append(item)
    if not results:
        raise ValidationError("업로드할 데이터 행이 없습니다.")

    # 2차 검증: 등번호를 이미 쓰는 기존 회원이 있으면, 그 회원도 이번 파일에서 번호가 바뀔 때만 허용한다.
    current = {r["id"]: r["back_number"] for r in conn.execute("SELECT id, back_number FROM member")}
    moving = {it["member_id"] for it in results
              if it["action"] == "갱신" and it["data"]["back_number"] != current.get(it["member_id"])}
    for it in results:
        if it["action"] == "오류" or it["data"]["back_number"] is None:
            continue
        owner = conn.execute("SELECT id, name FROM member WHERE back_number = ?", (it["data"]["back_number"],)).fetchone()
        if owner and owner["id"] != it["member_id"] and owner["id"] not in moving:
            it.update(action="오류", data=None, member_id=None,
                      reason=f"등번호 {it['data']['back_number']}번은 이미 {owner['name']} 회원이 사용 중입니다.")
    return results


def apply_members(conn, parsed: list[dict]) -> dict:
    """정상 행(신규·갱신)만 반영한다. 반영 중 실패한 행은 failed로 돌려준다."""
    done = {"신규": 0, "갱신": 0}
    failed = []
    ok = [p for p in parsed if p["action"] != "오류"]

    # 등번호가 바뀌는 기존 회원은 먼저 번호를 비워, 서로 맞바꾸는 경우에도 중복에 걸리지 않게 한다.
    original = {}
    for p in ok:
        if p["action"] == "갱신":
            cur = conn.execute("SELECT back_number FROM member WHERE id = ?", (p["member_id"],)).fetchone()
            if cur and cur["back_number"] is not None and cur["back_number"] != p["data"]["back_number"]:
                original[p["member_id"]] = cur["back_number"]
                conn.execute("UPDATE member SET back_number = NULL WHERE id = ?", (p["member_id"],))
    conn.commit()

    for item in ok:
        try:
            if item["action"] == "신규":
                create_member(conn, item["data"])
            else:
                update_member(conn, item["member_id"], item["data"])
                original.pop(item["member_id"], None)
            done[item["action"]] += 1
        except ValidationError as e:
            failed.append({"row": item["row"], "name": item["name"], "reason": " / ".join(e.messages)})
    # 반영에 실패한 회원은 원래 등번호로 되돌린다. 그 사이 다른 회원이 쓰게 된 번호는 비워 두고 알린다.
    for member_id, back in original.items():
        name = conn.execute("SELECT name FROM member WHERE id = ?", (member_id,)).fetchone()["name"]
        taken = conn.execute("SELECT 1 FROM member WHERE back_number = ?", (back,)).fetchone()
        if taken:
            failed.append({"row": "-", "name": name,
                           "reason": f"등번호 {back}번이 다른 회원에게 배정되어 {name} 회원의 등번호를 비워 두었습니다. 직접 지정하세요."})
        else:
            conn.execute("UPDATE member SET back_number = ? WHERE id = ?", (back, member_id))
    conn.commit()
    return {"신규": done["신규"], "갱신": done["갱신"], "실패": failed}
